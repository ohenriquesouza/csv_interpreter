import os
import re
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")

DEFAULT_THEME = "00AD32"

NULL_FILL = PatternFill(start_color="FFD7D7", end_color="FFD7D7", fill_type="solid")
ALERT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
THIN = Side(style="thin")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

_IMAGE_NAME_KEYWORDS = {"imagem", "foto", "image", "photo", "figura", "img", "picture", "thumb", "thumbnail"}
_IMAGE_EXT_RE = re.compile(r'\.(jpg|jpeg|png|gif|webp|bmp|tiff|svg)(\?.*)?$', re.IGNORECASE)


# --- Theme utilities ---

def _hex_to_rgb(hex_color: str) -> tuple[int, int, int]:
    h = hex_color.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _rgb_to_hex(r: int, g: int, b: int) -> str:
    return f"{r:02X}{g:02X}{b:02X}"


def _relative_luminance(r: int, g: int, b: int) -> float:
    def _ch(x: int) -> float:
        v = x / 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    return 0.2126 * _ch(r) + 0.7152 * _ch(g) + 0.0722 * _ch(b)


def _lighten(r: int, g: int, b: int, amount: float) -> tuple[int, int, int]:
    return (
        min(255, int(r + (255 - r) * amount)),
        min(255, int(g + (255 - g) * amount)),
        min(255, int(b + (255 - b) * amount)),
    )


def _darken(r: int, g: int, b: int, amount: float) -> tuple[int, int, int]:
    return (
        max(0, int(r * (1 - amount))),
        max(0, int(g * (1 - amount))),
        max(0, int(b * (1 - amount))),
    )


def build_theme(hex_color: str) -> dict:
    r, g, b = _hex_to_rgb(hex_color)
    lum = _relative_luminance(r, g, b)

    header_hex = _rgb_to_hex(r, g, b)

    lr, lg, lb = _lighten(r, g, b, 0.72)
    footer_hex = _rgb_to_hex(lr, lg, lb)

    header_text = "FFFFFF" if lum < 0.35 else "1A1A1A"

    dr, dg, db = _darken(r, g, b, 0.45)
    footer_text = _rgb_to_hex(dr, dg, db)

    return {
        "header_fill": PatternFill(start_color=header_hex, end_color=header_hex, fill_type="solid"),
        "footer_fill": PatternFill(start_color=footer_hex, end_color=footer_hex, fill_type="solid"),
        "header_text": header_text,
        "footer_text": footer_text,
    }


# --- Data helpers ---

def _is_null(value: str) -> bool:
    return str(value).strip().lower() in ("", "nan", "none", "null", "$null$", "n/a", "na")


def _is_image_col(col_name: str, series: pd.Series) -> bool:
    if any(kw in col_name.lower() for kw in _IMAGE_NAME_KEYWORDS):
        return True
    non_null = series[~series.apply(_is_null)]
    if len(non_null) == 0:
        return False
    return non_null.apply(lambda v: bool(_IMAGE_EXT_RE.search(str(v)))).mean() > 0.5


def _reorder_image_cols(df: pd.DataFrame) -> pd.DataFrame:
    image_cols = [col for col in df.columns if _is_image_col(col, df[col])]
    if not image_cols:
        return df
    other_cols = [col for col in df.columns if col not in image_cols]
    return df[other_cols + image_cols]


def _find_id_col(df: pd.DataFrame) -> str | None:
    for col in df.columns:
        if "id" in col.lower():
            return col
    return None


def _find_null_alerts(df: pd.DataFrame, threshold: int = 10) -> list[str]:
    id_col = _find_id_col(df)
    if id_col is None:
        return []

    alerts = []

    for unit_id, group in df.groupby(id_col, sort=False):
        cols_with_issues = []
        for col in df.columns:
            if col == id_col:
                continue
            null_count = group[col].apply(_is_null).sum()
            if null_count > threshold:
                cols_with_issues.append(f"{col} ({null_count})")
        if cols_with_issues:
            alerts.append(
                f"⚠️ {id_col} {unit_id} apresentou registros vazios em: "
                f"{', '.join(cols_with_issues)}. Possível erro no recebimento dos dados."
            )
    return alerts


def process_csv(
    csv_path: str,
    output_dir: str = OUTPUT_DIR,
    protect: bool = False,
    theme_color: str = DEFAULT_THEME,
    null_treatment: str = "na",
) -> str:
    df = pd.read_csv(csv_path, dtype=str, keep_default_na=False)
    df = _reorder_image_cols(df)
    theme = build_theme(theme_color)

    wb = Workbook()
    ws = wb.active
    ws.title = os.path.splitext(os.path.basename(csv_path))[0][:31]

    total_cols = len(df.columns)

    # --- Header row ---
    header_font = Font(bold=True, color=theme["header_text"], size=12)
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = theme["header_fill"]
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER
    ws.row_dimensions[1].height = 30

    # --- Data rows ---
    null_font = Font(color="999999", italic=True)
    null_count = 0
    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if _is_null(value):
                null_count += 1
                if null_treatment == "na":
                    cell.value = "N/A"
                    cell.fill = NULL_FILL
                    cell.font = null_font
                elif null_treatment == "zero":
                    cell.value = 0
                else:  # "vazio"
                    cell.value = ""
            else:
                cell.value = value
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = CELL_BORDER

    # --- Footer row ---
    footer_row = len(df) + 2

    def _write_footer_row(row_idx: int, text: str, fill: PatternFill, font: Font) -> None:
        ws.merge_cells(
            start_row=row_idx, start_column=1,
            end_row=row_idx, end_column=total_cols,
        )
        cell = ws.cell(row=row_idx, column=1)
        cell.value = text
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = CELL_BORDER
        ws.row_dimensions[row_idx].height = 25
        for col_idx in range(2, total_cols + 1):
            ws.cell(row=row_idx, column=col_idx).border = CELL_BORDER

    _write_footer_row(
        footer_row,
        (
            f"Total rows: {len(df)}  |  "
            f"Columns: {total_cols}  |  "
            f"Empty cells: {null_count}  |  "
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ),
        theme["footer_fill"],
        Font(bold=True, color=theme["footer_text"], size=11),
    )

    alerts = _find_null_alerts(df)
    for i, alert_msg in enumerate(alerts, start=1):
        _write_footer_row(
            footer_row + i,
            alert_msg,
            ALERT_FILL,
            Font(bold=True, color="7F6000", size=11),
        )

    # --- Column widths ---
    for col_idx, col_name in enumerate(df.columns, start=1):
        col_values = df.iloc[:, col_idx - 1].astype(str)
        max_len = max(len(str(col_name)), col_values.str.len().max()) + 4
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len, 50)

    ws.freeze_panes = "A2"

    if protect:
        ws.protection.sheet = True
        ws.protection.selectLockedCells = False

    os.makedirs(output_dir, exist_ok=True)
    filename = os.path.splitext(os.path.basename(csv_path))[0]
    out_path = os.path.join(output_dir, f"{filename}.xlsx")
    wb.save(out_path)
    return out_path


def main():
    if not os.path.isdir(DATA_DIR):
        print(f"[error] Data folder not found: {DATA_DIR}")
        return

    csv_files = [f for f in os.listdir(DATA_DIR) if f.lower().endswith(".csv")]
    if not csv_files:
        print("[info] No CSV files found in data/")
        return

    print(f"[info] Found {len(csv_files)} file(s)\n")
    for name in sorted(csv_files):
        path = os.path.join(DATA_DIR, name)
        print(f"  Processing: {name}")
        out = process_csv(path)
        print(f"  Saved:      {out}\n")

    print("[done]")


if __name__ == "__main__":
    main()
